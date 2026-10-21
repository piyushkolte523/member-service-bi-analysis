"""
Step 0 - Export the three data sheets from the assessment workbook to raw CSV files.

Why: DuckDB reads CSV natively (no extensions needed). The export is deliberately
"dumb": values are written exactly as stored in Excel - no trimming, no case fixes,
no de-duplication - so that every data-quality issue is still present when the
SQL scripts profile and clean the data.

Usage (from the project folder):
    python export_raw_csv.py
    python export_raw_csv.py --xlsx "path/to/BI_Analyst_Round_1_Technical_Assessment.xlsx"
"""
import argparse
import csv
import datetime as dt
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
DEFAULT_XLSX = HERE / "data" / "BI_Analyst_Round_1_Technical_Assessment.xlsx"
SHEETS = {"Members": "members_raw.csv",
          "Interactions": "interactions_raw.csv",
          "Targets": "targets_raw.csv"}


def fmt(v):
    """Write values faithfully: blanks stay blank, dates as ISO, numbers untouched."""
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    return v


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    ap.add_argument("--out", default=str(HERE / "data"))
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    for sheet, fname in SHEETS.items():
        ws = wb[sheet]
        n = 0
        with open(out_dir / fname, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                if all(v is None for v in row):
                    continue
                w.writerow([fmt(v) for v in row])
                n += 1
        print(f"{sheet:<13} -> {fname:<22} {n - 1:>5} data rows")


if __name__ == "__main__":
    main()
